"""CobrancaRelacaoExtractor: cobre as 3 variações GEOMATICS (CAMERAS, CAMERAS E
SENSORES, RASTREADORES DE CARGA). Todas seguem o mesmo padrão:

    - aba "COBRANÇA": tem dois valores "rotulados" (Valor por instalação /
      Valor mensalidade p/ veículo) que servem de fallback quando a linha não
      tem um valor próprio, e uma tabela PLACA + VALOR DA INSTALAÇÃO.
    - aba "RELAÇÃO DE VEÍCULOS COBRADOS": lista de placas que pagam a
      mensalidade (valor fixo, vindo do rótulo).

Duas dessas três variações (CAMERAS E SENSORES, RASTREADORES DE CARGA) também
aceitam a coluna se chamar "PLACAS" (plural) em vez de "PLACA" -- por isso o
parâmetro `aceitar_alias_placas`.
"""
from __future__ import annotations

from pathlib import Path
from typing import List

from openpyxl import load_workbook

from .. import utils
from ..config import conta_por_tipo_gasto
from ..modelos import Lancamento
from .base import ExtratorBase


class CobrancaRelacaoExtractor(ExtratorBase):
    def __init__(self, tipo_gasto: str, aceitar_alias_placas: bool = False):
        self.tipo_gasto = tipo_gasto
        self.aceitar_alias_placas = aceitar_alias_placas

    def _localizar_cabecalho(self, planilha, colunas_obrigatorias):
        for numero_linha, linha in enumerate(planilha.iter_rows(values_only=True), start=1):
            indices = {
                utils.normalizar_cabecalho(v): i
                for i, v in enumerate(linha)
                if v is not None
            }
            if self.aceitar_alias_placas and "placas" in indices and "placa" not in indices:
                indices["placa"] = indices["placas"]
            if colunas_obrigatorias.issubset(indices):
                return numero_linha, indices
        raise ValueError(f"As colunas {sorted(colunas_obrigatorias)} não foram encontradas na planilha '{planilha.title}'")

    def extrair(self, caminho_excel_fornecedor: Path) -> List[Lancamento]:
        conta = conta_por_tipo_gasto(self.tipo_gasto)

        wb = load_workbook(caminho_excel_fornecedor, read_only=False, data_only=True)
        try:
            planilha_cobranca = utils.localizar_planilha_por_nome(wb, "COBRANÇA")
            planilha_relacao = utils.localizar_planilha_por_nome(wb, "RELAÇÃO DE VEÍCULOS COBRADOS")

            valor_instalacao = utils.localizar_valor_rotulado(planilha_cobranca, "Valor por instalação")
            valor_mensalidade = utils.localizar_valor_rotulado(planilha_cobranca, "Valor mensalidade p/ veículo")

            linha_cab_inst, idx_inst = self._localizar_cabecalho(
                planilha_cobranca, {"placa", "valordainstalacao"}
            )

            lancamentos: List[Lancamento] = []

            for linha in planilha_cobranca.iter_rows(min_row=linha_cab_inst + 1, values_only=True):
                placa = linha[idx_inst["placa"]] if idx_inst["placa"] < len(linha) else None
                if placa is None:
                    continue
                valor = linha[idx_inst["valordainstalacao"]] if idx_inst["valordainstalacao"] < len(linha) else None
                lancamentos.append(
                    Lancamento(placa=placa, valor=valor if valor is not None else valor_instalacao, conta=conta)
                )

            linha_cab_rel, idx_rel = self._localizar_cabecalho(planilha_relacao, {"placa"})
            for linha in planilha_relacao.iter_rows(min_row=linha_cab_rel + 1, values_only=True):
                placa = linha[idx_rel["placa"]] if idx_rel["placa"] < len(linha) else None
                if placa is None:
                    continue
                lancamentos.append(Lancamento(placa=placa, valor=valor_mensalidade, conta=conta))

            return lancamentos
        finally:
            wb.close()
