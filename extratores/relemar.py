"""Dois extratores específicos da RELEMAR (RELEMAR MD usa o ColunaSimplesExtractor,
ver simples.py). Nenhum dos dois usa fallback de designação histórica -- por isso
`usar_historico = False` nas duas classes (igual ao comportamento original).
"""
from __future__ import annotations

from pathlib import Path
from typing import List

from openpyxl import load_workbook

from .. import utils
from ..config import conta_por_tipo_gasto
from ..modelos import Lancamento
from .base import ExtratorBase


class ListaAtivosExtractor(ExtratorBase):
    """RELEMAR GESTÃO ATIVOS: aba 'Lista de Ativos', colunas 'N° de placa' e
    'Valor Locação'. Placa tem o hífen removido."""

    usar_historico = False

    def __init__(self, tipo_gasto: str):
        self.tipo_gasto = tipo_gasto

    def extrair(self, caminho_excel_fornecedor: Path) -> List[Lancamento]:
        conta = conta_por_tipo_gasto(self.tipo_gasto)
        colunas_obrigatorias = {"ndeplaca", "valorlocacao"}

        wb = load_workbook(caminho_excel_fornecedor, read_only=True, data_only=True)
        try:
            planilha = utils.localizar_planilha_por_nome(wb, "Lista de Ativos")
            cabecalhos, linha_cabecalho = utils.encontrar_linha_cabecalho(planilha, colunas_obrigatorias)
            if cabecalhos is None:
                raise ValueError(
                    f"As colunas 'N° de placa' e 'Valor Locação' não foram encontradas em: "
                    f"{caminho_excel_fornecedor.name}"
                )
            idx = utils.indices_colunas(cabecalhos)

            lancamentos = []
            for linha in planilha.iter_rows(min_row=linha_cabecalho + 1, values_only=True):
                placa = linha[idx["ndeplaca"]] if idx["ndeplaca"] < len(linha) else None
                if placa is None:
                    continue
                placa = utils.limpar_placa(placa)
                valor = linha[idx["valorlocacao"]] if idx["valorlocacao"] < len(linha) else None
                lancamentos.append(Lancamento(placa=placa, valor=valor, conta=conta))
            return lancamentos
        finally:
            wb.close()


class MultasTaxaExtractor(ExtratorBase):
    """RELEMAR DESPACHANTE E MULTAS: aba 'MULTAS', cada linha gera DUAS
    lançamentos (valor da multa numa conta, taxa percentual em outra)."""

    usar_historico = False

    def __init__(self, tipo_gasto_multa: str, tipo_gasto_taxa: str):
        self.tipo_gasto_multa = tipo_gasto_multa
        self.tipo_gasto_taxa = tipo_gasto_taxa

    def extrair(self, caminho_excel_fornecedor: Path) -> List[Lancamento]:
        conta_multa = conta_por_tipo_gasto(self.tipo_gasto_multa)
        conta_taxa = conta_por_tipo_gasto(self.tipo_gasto_taxa)

        wb = load_workbook(caminho_excel_fornecedor, read_only=True, data_only=True)
        try:
            planilha = utils.localizar_planilha_por_nome(wb, "MULTAS")

            linha_cabecalho = None
            indices = None
            idx_placa = None

            for numero_linha, linha in enumerate(planilha.iter_rows(values_only=True), start=1):
                colunas = {
                    utils.normalizar_cabecalho(v): i
                    for i, v in enumerate(linha)
                    if v is not None
                }
                idx_placa_candidato = next((i for chave, i in colunas.items() if "placa" in chave), None)
                if idx_placa_candidato is not None and "multa" in colunas and "taxapercentual" in colunas:
                    linha_cabecalho = numero_linha
                    indices = colunas
                    idx_placa = idx_placa_candidato
                    break

            if linha_cabecalho is None:
                raise ValueError(
                    f"As colunas 'PLACA', 'MULTA' e 'TAXA PERCENTUAL' não foram encontradas em: "
                    f"{caminho_excel_fornecedor.name}"
                )

            lancamentos = []
            for linha in planilha.iter_rows(min_row=linha_cabecalho + 1, values_only=True):
                placa = linha[idx_placa]
                if placa is None:
                    continue
                placa = utils.limpar_placa(placa)

                valor_multa = linha[indices["multa"]]
                valor_taxa = linha[indices["taxapercentual"]]

                lancamentos.append(Lancamento(placa=placa, valor=valor_multa, conta=conta_multa))
                lancamentos.append(Lancamento(placa=placa, valor=valor_taxa, conta=conta_taxa))
            return lancamentos
        finally:
            wb.close()
