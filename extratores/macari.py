"""MultiSheetContaExtractor: cobre a MACARI, onde cada ABA do Excel do
fornecedor usa um tipo de gasto (conta) diferente -- ex.: aba 'NOTA FISCAL'
usa 'Serviços administrativos (externos)', aba 'REEMBOLSO' usa
'Licenças e placas'. Todas as abas mapeadas têm colunas PLACA + TOTAL.
Abas não mapeadas em `tipo_gasto_por_planilha` são ignoradas.
"""
from __future__ import annotations

from pathlib import Path
from typing import Dict, List

from openpyxl import load_workbook

from .. import utils
from ..config import conta_por_tipo_gasto
from ..modelos import Lancamento
from .base import ExtratorBase


class MultiSheetContaExtractor(ExtratorBase):
    def __init__(self, tipo_gasto_por_planilha: Dict[str, str]):
        # resolve a conta de cada aba já na construção (falha cedo se o tipo de
        # gasto não existir em GASTO_CONTAS)
        self._conta_por_planilha = {
            utils.normalizar_texto(nome_planilha): conta_por_tipo_gasto(tipo_gasto)
            for nome_planilha, tipo_gasto in tipo_gasto_por_planilha.items()
        }

    def extrair(self, caminho_excel_fornecedor: Path) -> List[Lancamento]:
        colunas_obrigatorias = {"placa", "total"}
        wb = load_workbook(caminho_excel_fornecedor, read_only=True, data_only=True)
        lancamentos: List[Lancamento] = []
        try:
            for planilha in wb.worksheets:
                conta = self._conta_por_planilha.get(utils.normalizar_texto(planilha.title))
                if conta is None:
                    continue  # aba fora do mapeamento configurado para este fornecedor

                cabecalhos, linha_cabecalho = utils.encontrar_linha_cabecalho(planilha, colunas_obrigatorias)
                if cabecalhos is None:
                    raise ValueError(
                        f"As colunas 'PLACA' e 'TOTAL' não foram encontradas na aba "
                        f"'{planilha.title}' de: {caminho_excel_fornecedor.name}"
                    )
                idx = utils.indices_colunas(cabecalhos)

                for linha in planilha.iter_rows(min_row=linha_cabecalho + 1, values_only=True):
                    placa = linha[idx["placa"]] if idx["placa"] < len(linha) else None
                    if placa is None:
                        continue
                    valor = linha[idx["total"]] if idx["total"] < len(linha) else None
                    lancamentos.append(Lancamento(placa=placa, valor=valor, conta=conta))
        finally:
            wb.close()
        return lancamentos
