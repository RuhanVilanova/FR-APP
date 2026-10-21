"""TipoTransacaoExtractor: cobre a EDENRED PEDÁGIO. É o extrator mais
complexo porque:

  1) a conta/serviço de cada linha vem do "Tipo de Transação" (mapeado em
     REGRAS_TIPO_TRANSACAO), com um relatório à parte de tipos não mapeados;
  2) depois de enriquecido com a designação, é preciso ABATER cada linha de
     'Estorno' (valor positivo) contra linha(s) de 'Pedágio' (valor negativo)
     do MESMO grupo AN/CCO/CC -- isso só pode acontecer depois do
     enriquecimento, por isso vira o hook `pos_processar`.
"""
from __future__ import annotations

from collections import defaultdict
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook

from .. import utils
from ..config import conta_por_tipo_gasto
from ..modelos import Lancamento
from .base import ExtratorBase

# tipo de transação (normalizado) -> (tipo de gasto | None, texto p/ coluna SERVIÇO)
# tipo de gasto None = conta fica vazia (caso do Estorno)
REGRAS_TIPO_TRANSACAO: Dict[str, Tuple[Optional[str], str]] = {
    "pedagio": ("Pedágios", "Pedágio"),
    "estorno pedagio": (None, "Estorno"),
    "estacionamento": ("Gastos com viagem Estacionamento", "Estacionamento"),
    "mensalidade tag ativa - licenca de software": ("Serviços administrativos (externos)", "Mensalidade"),
}


def _chave_grupo(an, cco, cc) -> tuple:
    return (utils.normalizar_texto(an), utils.normalizar_texto(cco), utils.normalizar_texto(cc))


class TipoTransacaoExtractor(ExtratorBase):
    incluir_servico = True

    def __init__(self, regras_tipo_transacao: Dict[str, Tuple[Optional[str], str]] = REGRAS_TIPO_TRANSACAO):
        self.regras_tipo_transacao = regras_tipo_transacao
        # preenchido durante extrair(); processor/relatório final leem depois
        self.tipos_nao_mapeados: List[dict] = []

    def _resolver_conta_e_servico(self, tipo_transacao) -> Tuple[Optional[str], str, bool]:
        regra = self.regras_tipo_transacao.get(utils.normalizar_espacos(tipo_transacao))
        if regra is None:
            return None, "NÃO MAPEADO", False
        tipo_gasto, servico = regra
        if tipo_gasto is None:
            return None, servico, True
        return conta_por_tipo_gasto(tipo_gasto), servico, True

    def extrair(self, caminho_excel_fornecedor: Path) -> List[Lancamento]:
        self.tipos_nao_mapeados = []

        chave_placa = utils.normalizar_cabecalho("Placa")
        chave_valor = utils.normalizar_cabecalho("Valor cobrado(R$)")
        chave_tipo_transacao = utils.normalizar_cabecalho("Tipo de Transação")
        colunas_obrigatorias = {chave_placa, chave_valor, chave_tipo_transacao}

        wb = load_workbook(caminho_excel_fornecedor, read_only=True, data_only=True)
        lancamentos: List[Lancamento] = []
        abas_processadas = 0
        try:
            for planilha in wb.worksheets:
                cabecalhos, linha_cabecalho = utils.encontrar_linha_cabecalho(planilha, colunas_obrigatorias)
                if cabecalhos is None:
                    continue
                abas_processadas += 1
                idx = utils.indices_colunas(cabecalhos)

                for linha in planilha.iter_rows(min_row=linha_cabecalho + 1, values_only=True):
                    placa = linha[idx[chave_placa]] if idx[chave_placa] < len(linha) else None
                    if placa is None:
                        continue

                    valor = linha[idx[chave_valor]] if idx[chave_valor] < len(linha) else None
                    tipo_transacao = linha[idx[chave_tipo_transacao]] if idx[chave_tipo_transacao] < len(linha) else None

                    conta, servico, mapeado = self._resolver_conta_e_servico(tipo_transacao)
                    if not mapeado:
                        self.tipos_nao_mapeados.append({"placa": placa, "tipo_transacao": tipo_transacao})

                    lancamentos.append(Lancamento(placa=placa, valor=valor, conta=conta, servico=servico))
        finally:
            wb.close()

        if abas_processadas == 0:
            raise ValueError(
                f"Nenhuma aba com as colunas 'Placa', 'Valor cobrado(R$)' e 'Tipo de Transação' "
                f"foi encontrada em: {caminho_excel_fornecedor.name}"
            )
        return lancamentos

    def pos_processar(self, lancamentos: List[Lancamento]) -> List[Lancamento]:
        """Abate cada 'Estorno' (valor positivo) contra 'Pedágio' (valor negativo)
        do mesmo grupo AN/CCO/CC, na ordem em que aparecem:
          - Pedágio totalmente consumido -> linha excluída.
          - Pedágio parcialmente consumido -> valor reduzido, linha mantida.
          - Estorno totalmente abatido -> linha excluída.
          - Estorno sem Pedágio suficiente -> linha mantida com o valor restante.
        No final, todos os valores voltam a ser positivos (abs).
        """
        grupos_pedagio = defaultdict(list)
        for i, lanc in enumerate(lancamentos):
            if lanc.servico == "Pedágio":
                grupos_pedagio[_chave_grupo(lanc.an, lanc.cco, lanc.cc)].append(i)

        removidos = set()

        for i_estorno, lanc_estorno in enumerate(lancamentos):
            if lanc_estorno.servico != "Estorno" or lanc_estorno.valor is None:
                continue

            chave = _chave_grupo(lanc_estorno.an, lanc_estorno.cco, lanc_estorno.cc)
            restante = abs(lanc_estorno.valor)

            for i_pedagio in grupos_pedagio.get(chave, []):
                if restante <= 0:
                    break
                if i_pedagio in removidos:
                    continue

                lanc_pedagio = lancamentos[i_pedagio]
                if lanc_pedagio.valor is None:
                    continue
                valor_pedagio_abs = abs(lanc_pedagio.valor)

                if valor_pedagio_abs <= restante:
                    restante -= valor_pedagio_abs
                    removidos.add(i_pedagio)
                else:
                    lanc_pedagio.valor = lanc_pedagio.valor + restante
                    restante = 0

            if restante <= 0:
                removidos.add(i_estorno)
            else:
                lanc_estorno.valor = restante

        resultado = [lanc for i, lanc in enumerate(lancamentos) if i not in removidos]
        for lanc in resultado:
            if lanc.valor is not None:
                lanc.valor = abs(lanc.valor)
        return resultado
