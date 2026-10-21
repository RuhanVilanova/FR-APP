"""DescricaoMapeadaExtractor: cobre a UNIDAS LEVES E PESADOS.

Aqui a conta/serviço de cada linha não vem de um "tipo de gasto fixo do
fornecedor" (como VAMOS, TKS etc.), e sim da coluna "Descrição" de cada linha,
via um dicionário de mapeamento. Além disso, para a descrição
"LOCACAO VEICULOS PESADOS" a placa correta vem de uma coluna diferente
("Placa correta") em vez da coluna "Placa do veículo".

Percorre TODAS as abas do arquivo (as colunas se repetem em várias abas).
"""
from __future__ import annotations

from pathlib import Path
from typing import Dict, List, Set, Tuple

from openpyxl import load_workbook

from .. import utils
from ..config import conta_por_tipo_gasto
from ..modelos import Lancamento
from .base import ExtratorBase

# Descrição (normalizada) -> tipo de gasto (usado para achar a conta)
DESCRICAO_PARA_TIPO_GASTO: Dict[str, str] = {
    "reembolso multa leves": "Multas estaduais por penalidades",
    "reembolso multas de transito": "Multas estaduais por penalidades",
    "taxa administrativa leves": "Serviços administrativos (externos)",
    "locacao veiculos leves": "Locação de veículos",
    "locacao veiculos pesados": "Locação de veículos",
    "reaj retroativo 25%": "Locação de veículos",
}

# Descrição (normalizada) -> texto da coluna SERVIÇO no CONSO
DESCRICAO_PARA_SERVICO: Dict[str, str] = {
    "reembolso multa leves": "Reembolso multa leves",
    "reembolso multas de transito": "Reembolso multa pesados",
    "taxa administrativa leves": "Reembolso multa leves",
    "locacao veiculos leves": "Locação  Leves",
    "locacao veiculos pesados": "Locação Pesados",
    "reaj retroativo 25%": "Rejuste",
}

# Descrições cuja placa deve vir da coluna "Placa correta"
DESCRICOES_USAM_PLACA_CORRETA: Set[str] = {"locacao veiculos pesados"}


class DescricaoMapeadaExtractor(ExtratorBase):
    incluir_servico = True

    def __init__(
        self,
        descricao_para_tipo_gasto: Dict[str, str] = DESCRICAO_PARA_TIPO_GASTO,
        descricao_para_servico: Dict[str, str] = DESCRICAO_PARA_SERVICO,
        descricoes_usam_placa_correta: Set[str] = DESCRICOES_USAM_PLACA_CORRETA,
    ):
        self.descricao_para_tipo_gasto = descricao_para_tipo_gasto
        self.descricao_para_servico = descricao_para_servico
        self.descricoes_usam_placa_correta = descricoes_usam_placa_correta

    def extrair(self, caminho_excel_fornecedor: Path) -> List[Lancamento]:
        wb = load_workbook(caminho_excel_fornecedor, read_only=True, data_only=True)
        lancamentos: List[Lancamento] = []
        try:
            for planilha in wb.worksheets:
                cabecalhos = None
                linha_cabecalho = None
                indices = None

                for numero_linha, linha in enumerate(planilha.iter_rows(values_only=True), start=1):
                    colunas = {
                        utils.normalizar_cabecalho(v): i
                        for i, v in enumerate(linha)
                        if v is not None
                    }
                    if {"descricao", "placadoveiculo", "montitemcalculavel"}.issubset(colunas):
                        cabecalhos = linha
                        linha_cabecalho = numero_linha
                        indices = colunas
                        break

                if cabecalhos is None:
                    continue  # aba sem as colunas esperadas

                idx_descricao = indices["descricao"]
                idx_placa = indices["placadoveiculo"]
                idx_valor = indices["montitemcalculavel"]
                idx_placa_correta = indices.get("placacorreta")

                for linha in planilha.iter_rows(min_row=linha_cabecalho + 1, values_only=True):
                    descricao = linha[idx_descricao] if idx_descricao < len(linha) else None
                    if descricao is None:
                        continue
                    descricao_norm = utils.normalizar_texto(descricao)

                    tipo_gasto = self.descricao_para_tipo_gasto.get(descricao_norm)
                    if tipo_gasto is None:
                        continue  # descrição não mapeada: ignora a linha (igual ao original)

                    conta = conta_por_tipo_gasto(tipo_gasto)
                    servico = self.descricao_para_servico.get(descricao_norm)

                    if descricao_norm in self.descricoes_usam_placa_correta and idx_placa_correta is not None:
                        placa = linha[idx_placa_correta] if idx_placa_correta < len(linha) else None
                    else:
                        placa = linha[idx_placa] if idx_placa < len(linha) else None
                    if placa is None:
                        continue
                    placa = utils.limpar_placa(placa)

                    valor = linha[idx_valor] if idx_valor < len(linha) else None
                    lancamentos.append(Lancamento(placa=placa, valor=valor, conta=conta, servico=servico))
        finally:
            wb.close()
        return lancamentos
