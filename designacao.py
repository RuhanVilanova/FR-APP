"""DesignacaoService: centraliza a leitura dos Excels de "Designação" (normal e histórico)
e a busca PLACA -> (AN, CCO, CC, TIPO).

No código original essa lógica (carregarDadosDesignacao / carregarDesignacaoHistorica /
buscarDesignacaoHistorica) estava duplicada, praticamente idêntica, em quase todos os
scripts (exceto RELEMAR, que nunca chamava o fallback histórico). Aqui vira um único
serviço reutilizável, com cache do histórico (ele é o mesmo arquivo para todos os
fornecedores, então só precisa ser lido uma vez por execução).
"""
from __future__ import annotations

from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

from . import utils

DadosDesignacao = tuple  # (an, cco, cc, tipo)

_COLUNAS_DESIGNACAO = {"placa", "an", "cco", "cc", "tipo"}


class DesignacaoService:
    """Uso típico:

        servico = DesignacaoService(caminho_designacao_historica)
        servico.carregar_designacao_do_fornecedor(caminho_excel_fornecedor)
        an, cco, cc, tipo = servico.buscar(placa, usar_historico=True)
    """

    # cache de classe: o histórico é o MESMO arquivo para todos os fornecedores
    _cache_historico: Optional[dict[str, DadosDesignacao]] = None

    def __init__(self, caminho_designacao_historica: Path):
        self._caminho_designacao_historica = caminho_designacao_historica
        self._dados_fornecedor: dict[str, DadosDesignacao] = {}

    # -- designação "normal" (específica do fornecedor / mês) ----------------
    def carregar_designacao_do_fornecedor(self, caminho_excel_fornecedor: Path) -> None:
        caminho_excel_designacao = utils.pegar_excel_designacao(caminho_excel_fornecedor)
        self._dados_fornecedor = self._carregar_de_tabelas(caminho_excel_designacao)

    @staticmethod
    def _carregar_de_tabelas(caminho_excel_designacao: Path) -> dict[str, DadosDesignacao]:
        """Lê a primeira aba (priorizando a aba 'BASE') que tenha uma Tabela do Excel
        com as colunas PLACA/AN/CCO/CC/TIPO."""
        wb = load_workbook(caminho_excel_designacao, read_only=False, data_only=True)
        planilhas_priorizadas = sorted(
            wb.worksheets,
            key=lambda p: (utils.normalizar_texto(p.title) != "base", utils.normalizar_texto(p.title)),
        )

        for planilha in planilhas_priorizadas:
            for tabela in planilha.tables.values():
                colunas = {
                    utils.normalizar_cabecalho(c.name): i
                    for i, c in enumerate(tabela.tableColumns)
                }
                if not _COLUNAS_DESIGNACAO.issubset(colunas):
                    continue

                dados = {}
                for linha in planilha[tabela.ref][1:]:
                    valores = [c.value for c in linha]
                    placa = valores[colunas["placa"]] if colunas["placa"] < len(valores) else None
                    if placa is None:
                        continue
                    dados[utils.normalizar_texto(placa)] = (
                        valores[colunas["an"]] if colunas["an"] < len(valores) else None,
                        valores[colunas["cco"]] if colunas["cco"] < len(valores) else None,
                        valores[colunas["cc"]] if colunas["cc"] < len(valores) else None,
                        valores[colunas["tipo"]] if colunas["tipo"] < len(valores) else None,
                    )
                wb.close()
                return dados

        wb.close()
        raise ValueError(
            f"Nenhuma tabela com as colunas 'PLACA', 'AN', 'CCO', 'CC' e 'TIPO' foi "
            f"encontrada em: {caminho_excel_designacao.name}"
        )

    # -- designação histórica (arquivo único em AUXILIARES) ------------------
    def _carregar_historico(self) -> dict[str, DadosDesignacao]:
        if DesignacaoService._cache_historico is not None:
            return DesignacaoService._cache_historico

        if not self._caminho_designacao_historica.exists():
            raise FileNotFoundError(
                f"Arquivo de designação histórica não encontrado: {self._caminho_designacao_historica}"
            )

        wb = load_workbook(self._caminho_designacao_historica, read_only=True, data_only=True)
        dados_historico: dict[str, DadosDesignacao] = {}

        for planilha in wb.worksheets:
            cabecalhos, linha_cabecalho = utils.encontrar_linha_cabecalho(planilha, _COLUNAS_DESIGNACAO)
            if cabecalhos is None:
                continue

            idx = utils.indices_colunas(cabecalhos)
            for linha in planilha.iter_rows(min_row=linha_cabecalho + 1, values_only=True):
                placa = linha[idx["placa"]] if idx["placa"] < len(linha) else None
                if placa is None:
                    continue
                dados_historico[utils.normalizar_texto(placa)] = (
                    linha[idx["an"]] if idx["an"] < len(linha) else None,
                    linha[idx["cco"]] if idx["cco"] < len(linha) else None,
                    linha[idx["cc"]] if idx["cc"] < len(linha) else None,
                    linha[idx["tipo"]] if idx["tipo"] < len(linha) else None,
                )

        wb.close()
        DesignacaoService._cache_historico = dados_historico
        return dados_historico

    def _buscar_historico(self, placa) -> Optional[DadosDesignacao]:
        """1) Placa exata no histórico. 2) Qualquer placa do histórico com o mesmo
        prefixo de 3 caracteres. 3) None."""
        historico = self._carregar_historico()
        placa_norm = utils.normalizar_texto(placa)

        dados = historico.get(placa_norm)
        if dados is not None:
            return dados

        prefixo = placa_norm[:3]
        if not prefixo:
            return None

        for placa_hist, dados in historico.items():
            if placa_hist[:3] == prefixo:
                return dados
        return None

    # -- API pública -----------------------------------------------------------
    def buscar(self, placa, usar_historico: bool = True) -> DadosDesignacao:
        """Retorna (an, cco, cc, tipo). Se não achar nada, retorna (None, None, None, None)."""
        dados = self._dados_fornecedor.get(utils.normalizar_texto(placa))
        if dados is None and usar_historico:
            dados = self._buscar_historico(placa)
        return dados or (None, None, None, None)